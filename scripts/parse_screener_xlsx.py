"""Parse screener.in export xlsx 'Data Sheet' -> annual + quarterly earnings.

The export's deep history is ANNUAL (~10yr Profit&Loss); quarterly stays ~10q.
We extract annual Net Profit + Sales (and quarterly) and compute YoY growth.
Net-profit growth is used as the earnings-growth metric (avoids messy
EPS-from-shares with splits/bonuses).
Output: data/reference/screener_annual.csv (symbol, period_end, freq, net_profit, sales, yoy_growth)
"""
from __future__ import annotations
import sys
from datetime import date
from pathlib import Path
import openpyxl, polars as pl

_ROOT = Path(__file__).resolve().parents[1]
XLSX_DIR = _ROOT / "data" / "raw" / "screener_xlsx"
OUT = _ROOT / "data" / "reference" / "screener_annual.csv"

def _row_label(ws, r): 
    v = ws.cell(r,1).value
    return str(v).strip().lower() if v else ""

def _section_rows(ws, header_kw):
    """Return (date_row, {label: row}) for the section starting at header_kw."""
    start=None
    for r in range(1, ws.max_row+1):
        if _row_label(ws,r).startswith(header_kw.lower()):
            start=r; break
    if start is None: return None,None
    date_row=None; labels={}
    for r in range(start+1, min(start+40, ws.max_row+1)):
        lab=_row_label(ws,r)
        if lab.startswith("report da"): date_row=r
        elif lab in ("sales","net profit"): labels[lab]=r
        elif lab and lab.upper()==lab and r>start+2 and not any(c in lab for c in "0123456789"):
            pass
        # stop at next big section header
        if r>start+1 and _row_label(ws,r) in ("quarters","balance sheet","cash flow","derived:"):
            break
    return date_row, labels

def parse_file(path: Path, symbol: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Data Sheet"]
    rows=[]
    for header,freq in [("profit & loss","annual"),("quarters","quarterly")]:
        dr, labs = _section_rows(ws, header)
        if dr is None or "net profit" not in labs: continue
        dates=[ws.cell(dr,c).value for c in range(2, ws.max_column+1)]
        npr=labs.get("net profit"); sr=labs.get("sales")
        for i,dt in enumerate(dates):
            if dt is None: continue
            if hasattr(dt,"date"): dt=dt.date()
            rows.append({"symbol":symbol,"period_end":dt,"freq":freq,
                         "net_profit": ws.cell(npr,2+i).value,
                         "sales": ws.cell(sr,2+i).value if sr else None})
    return rows

def main():
    files=sorted(XLSX_DIR.glob("*.xlsx"))
    if len(sys.argv)>1: files=[XLSX_DIR/f"{sys.argv[1]}.xlsx"]
    allrows=[]
    for f in files:
        try: allrows+=parse_file(f, f.stem)
        except Exception as e: print(f"  {f.stem}: parse err {e}")
    df=pl.DataFrame(allrows)
    if df.is_empty(): print("no data"); return
    # YoY growth within freq: annual shift 1, quarterly shift 4
    out=[]
    for freq,shift in [("annual",1),("quarterly",4)]:
        d=df.filter(pl.col("freq")==freq).sort(["symbol","period_end"])
        d=d.with_columns((pl.col("net_profit")/pl.col("net_profit").shift(shift).over("symbol")-1).alias("yoy_growth"))
        out.append(d)
    res=pl.concat(out).sort(["symbol","freq","period_end"])
    OUT.parent.mkdir(parents=True,exist_ok=True); res.write_csv(OUT)
    ann=res.filter((pl.col("freq")=="annual")&pl.col("yoy_growth").is_not_null())
    print(f"parsed {df['symbol'].n_unique()} symbols, {res.height} rows")
    print(f"annual rows with YoY: {ann.height}, date range {ann['period_end'].min()}..{ann['period_end'].max()}")
    print(res.filter(pl.col("symbol")==files[0].stem).filter(pl.col("freq")=="annual").tail(6))

if __name__=="__main__": main()
